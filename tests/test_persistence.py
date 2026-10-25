import csv
import json
import os
import sqlite3
from dataclasses import replace
from datetime import datetime, timedelta, timezone

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np
import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from app.config import AudioConfig, CalibrationProfile, MeasurementConfig
from core.latency_analyzer import LatencyMeasurement
from core.persistence_controller import PersistenceController
from core.persistence_service import PersistenceService
from core.test_models import MeasurementRecord
from core.test_models import TestSessionConfig as SessionConfig
from core.test_models import TestSessionSnapshot as SessionSnapshot
from core.test_state_machine import TestState as SessionState
from data.database import Database, DatabaseError
from data.repository import DuplicateSessionError, SessionNotFoundError, SessionRepository
from ui.history_page import HistoryPage


def _snapshot(
    session_id: str = "session-1",
    headphone_name: str = "Alpha",
    connection_type: str = "bluetooth",
    with_calibration: bool = False,
    latency_offset_ms: float = 0.0,
) -> SessionSnapshot:
    started = datetime(2026, 10, 25, 12, 0, tzinfo=timezone.utc)
    audio = AudioConfig(
        backend="simulation",
        sample_rate=48000,
        output_device_id="sim-output",
        input_device_id="sim-input",
    )
    measurement_config = MeasurementConfig(
        default_test_count=3,
        minimum_valid_count=1,
        warmup_seconds=1.0,
        inter_measurement_delay_ms=0.0,
    )
    profile = None
    if with_calibration:
        profile = CalibrationProfile(
            profile_id="profile-1",
            name="环回校准",
            output_device_id="sim-output",
            output_device="模拟输出",
            input_device_id="sim-input",
            input_device="模拟输入",
            sample_rate=48000,
            audio_mode="simulation",
            compensation_ms=2.0,
            calibration_method="合成已知延迟",
            calibration_date="2026-10-25T11:00:00+00:00",
            notes="仅用于持久化追溯",
        )
    config = SessionConfig(
        audio=audio,
        measurement=measurement_config,
        compensation_ms=2.0 if with_calibration else 0.0,
        calibration_profile=profile,
        output_device_name="模拟输出",
        input_device_name="模拟输入",
        started_at=started,
        session_id=session_id,
        session_name=f"{headphone_name} 测试",
        headphone_name=headphone_name,
        connection_type=connection_type,
        notes="初始备注",
    )
    records = []
    samples = (
        (20.0 + latency_offset_ms, True),
        (30.0 + latency_offset_ms, True),
        (250.0 + latency_offset_ms, False),
    )
    for sequence, (raw, valid) in enumerate(samples, start=1):
        timestamp = started + timedelta(seconds=sequence)
        result = LatencyMeasurement(
            raw_latency_ms=raw,
            compensated_latency_ms=raw - config.compensation_ms,
            compensation_ms=config.compensation_ms,
            correlation_lag_samples=480.0 + sequence,
            correlation_peak=0.92 if valid else 0.2,
            correlation_ambiguity_ratio=0.1,
            signal_amplitude=0.2,
            noise_floor=0.01,
            snr_db=20.0 if valid else float("inf"),
            is_valid=valid,
            invalid_reason=None if valid else "latency_out_of_range",
            timestamp=timestamp,
            playback_timestamp_seconds=100.0,
            capture_timestamp_seconds=100.0 + raw / 1000,
            is_simulation=True,
            timing_basis="synthetic_sample_index",
            timing_limitations=("synthetic_time_axis",),
        )
        records.append(
            MeasurementRecord(
                sequence=sequence,
                reference=np.asarray([0.0, 0.5, -0.5, 0.0], dtype=np.float32),
                captured=np.asarray([0.0, 0.0, 0.5, -0.5, 0.0], dtype=np.float32),
                capture_window_start_sample=0,
                playback_sample_index=0,
                measurement=result,
            )
        )
    return SessionSnapshot(
        config=config,
        records=tuple(records),
        terminal_state=SessionState.COMPLETED,
        started_at=started,
        finished_at=started + timedelta(minutes=1),
        is_simulation=True,
    )


def test_schema_save_reopen_and_recalculate_statistics(tmp_path):
    path = tmp_path / "history.db"
    service = PersistenceService(path)
    snapshot = _snapshot(with_calibration=True)

    session_id = service.save_session(snapshot)
    database = Database(path)
    database.initialize()
    assert database.schema_version() == 1

    reopened = PersistenceService(path)
    loaded = reopened.get_session(session_id)
    assert loaded.session_uid == "session-1"
    assert loaded.headphone_name == "Alpha"
    assert loaded.connection_type == "bluetooth"
    assert loaded.completed_count == loaded.valid_count + loaded.invalid_count == 3
    assert loaded.valid_count == 2
    assert loaded.invalid_count == 1
    assert loaded.calibration_profile is not None
    assert loaded.calibration_profile.profile_id == "profile-1"
    assert loaded.calibration_profile.notes == "仅用于持久化追溯"
    assert json.loads(loaded.audio_config_json)["sample_rate"] == 48000
    assert len(loaded.results) == 3
    assert loaded.results[-1].is_valid is False
    assert loaded.results[-1].invalid_reason == "latency_out_of_range"
    assert np.frombuffer(loaded.results[0].reference_waveform, dtype="<f4").tolist() == [
        0.0,
        0.5,
        -0.5,
        0.0,
    ]
    recalculated = reopened.recalculate_statistics(loaded)
    assert recalculated.valid_count == snapshot.statistics.valid_count
    assert recalculated.invalid_count == snapshot.statistics.invalid_count
    assert recalculated.raw.mean_latency_ms == pytest.approx(25.0)
    assert recalculated.compensated.mean_latency_ms == pytest.approx(23.0)
    assert recalculated.raw.p95_latency_ms == pytest.approx(snapshot.statistics.raw.p95_latency_ms)
    assert reopened.list_sessions()[0].session_uid == "session-1"


def test_duplicate_save_is_rejected_without_duplicate_rows(tmp_path):
    service = PersistenceService(tmp_path / "history.db")
    snapshot = _snapshot()
    service.save_session(snapshot)

    with pytest.raises(DuplicateSessionError):
        service.save_session(snapshot)

    sessions = service.list_sessions()
    assert len(sessions) == 1
    assert len(service.get_session(sessions[0].id).results) == 3


def test_history_search_filter_sort_notes_and_cascade_delete(tmp_path):
    path = tmp_path / "history.db"
    service = PersistenceService(path)
    first_id = service.save_session(
        _snapshot("first", "Alpha", "bluetooth", with_calibration=True)
    )
    second_snapshot = _snapshot("second", "Beta", "wired", latency_offset_ms=10.0)
    second_snapshot = replace(
        second_snapshot,
        config=replace(second_snapshot.config, session_name="Beta 测试"),
    )
    second_id = service.save_session(second_snapshot)

    assert [session.id for session in service.list_sessions(search="Alpha")] == [first_id]
    assert [session.id for session in service.list_sessions(connection_type="wired")] == [second_id]
    assert [session.id for session in service.list_sessions(sort_by="mean_latency_asc")] == [
        first_id,
        second_id,
    ]
    service.update_notes(first_id, "  修改后的备注  ")
    assert service.get_session(first_id).notes == "修改后的备注"

    assert service.delete_sessions([first_id]) == 1
    with pytest.raises(SessionNotFoundError):
        service.get_session(first_id)
    with sqlite3.connect(path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM test_results WHERE session_id = ?", (first_id,)).fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM calibration_profiles WHERE session_id = ?", (first_id,)).fetchone()[0] == 0
    assert [session.id for session in service.list_sessions()] == [second_id]


def test_failed_save_rolls_back_session_and_calibration(tmp_path, monkeypatch):
    database = Database(tmp_path / "history.db")
    repository = SessionRepository(database)
    snapshot = _snapshot(with_calibration=True)
    stored = PersistenceService._to_stored_session(snapshot)

    def fail_insert(*args):
        raise sqlite3.OperationalError("模拟写入失败")

    monkeypatch.setattr(repository, "_insert_calibration", fail_insert)
    with pytest.raises(sqlite3.OperationalError):
        repository.save_session(stored)
    assert repository.list_sessions() == []


def test_newer_schema_is_reported(tmp_path):
    path = tmp_path / "history.db"
    with sqlite3.connect(path) as connection:
        connection.execute("PRAGMA user_version = 99")
    with pytest.raises(DatabaseError):
        PersistenceService(path).list_sessions()


def test_csv_and_excel_export_match_session_and_result_fields(tmp_path):
    service = PersistenceService(tmp_path / "history.db")
    session_id = service.save_session(_snapshot())
    csv_path = service.export_sessions([session_id], tmp_path / "records.csv", "csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 3
    assert rows[0]["session_uid"] == "session-1"
    assert rows[0]["headphone_name"] == "Alpha"
    assert rows[0]["raw_latency_ms"] == "20.0"
    assert rows[0]["compensated_latency_ms"] == "20.0"
    assert rows[0]["is_simulation"] == "True"
    assert rows[-1]["invalid_reason"] == "latency_out_of_range"
    assert rows[-1]["snr_db"] == "inf"
    assert rows[0]["reference_sample_count"] == "4"

    excel_path = service.export_sessions([session_id], tmp_path / "records.xlsx", "xlsx")
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["会话", "单次测量"]
    session_rows = list(workbook["会话"].iter_rows(values_only=True))
    result_rows = list(workbook["单次测量"].iter_rows(values_only=True))
    assert len(session_rows) == 2
    assert len(result_rows) == 4
    session_headers = session_rows[0]
    result_headers = result_rows[0]
    assert session_rows[1][session_headers.index("session_uid")] == "session-1"
    assert result_rows[1][result_headers.index("raw_latency_ms")] == 20.0
    assert result_rows[-1][result_headers.index("snr_db")] == "inf"
    workbook.close()


@pytest.fixture
def qt_application():
    application = QApplication.instance() or QApplication([])
    yield application


@pytest.fixture
def history_context(qt_application, tmp_path):
    service = PersistenceService(tmp_path / "history-ui.db")
    controller = PersistenceController(service)
    page = HistoryPage(controller)
    page.show()
    yield service, controller, page
    page._search_timer.stop()
    controller.shutdown()
    page.close()
    page.deleteLater()
    controller.deleteLater()
    qt_application.processEvents()


def _wait_for_signal(signal, timeout_ms=3000):
    event_loop = QEventLoop()
    timer = QTimer()
    timer.setSingleShot(True)
    timer.timeout.connect(event_loop.quit)
    signal.connect(event_loop.quit)
    timer.start(timeout_ms)
    event_loop.exec()
    signal.disconnect(event_loop.quit)
    received = timer.isActive()
    timer.stop()
    return received


def _session_row(page, session_id):
    return next(row for row, session in enumerate(page._sessions) if session.id == session_id)


def test_history_page_search_filter_and_sort(history_context):
    service, controller, page = history_context
    first_id = service.save_session(_snapshot("first", "Alpha", "bluetooth"))
    second_id = service.save_session(
        _snapshot("second", "Beta", "wired", latency_offset_ms=10.0)
    )

    page.refresh()
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert [session.id for session in page._sessions] == [second_id, first_id]
    assert page.table.rowCount() == 2

    page.sort_combo.setCurrentIndex(page.sort_combo.findData("mean_latency_asc"))
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert [session.id for session in page._sessions] == [first_id, second_id]
    assert page.table.item(0, 4).text() == "25.00 ms"
    assert page.table.item(1, 4).text() == "35.00 ms"

    page.search_input.setText("Alpha")
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert [session.id for session in page._sessions] == [first_id]
    assert page.table.rowCount() == 1

    page.search_input.clear()
    assert _wait_for_signal(controller.sessions_loaded) is True
    page.connection_filter.setCurrentIndex(page.connection_filter.findData("wired"))
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert [session.id for session in page._sessions] == [second_id]
    assert page.table.item(0, 1).text() == "Beta"
    assert page.table.item(0, 2).text() == "有线"


def test_history_page_details_notes_and_multi_selection_restore(history_context):
    service, controller, page = history_context
    first_id = service.save_session(_snapshot("first", "Alpha", "bluetooth"))
    second_id = service.save_session(
        _snapshot("second", "Beta", "wired", latency_offset_ms=10.0)
    )
    page.refresh()
    assert _wait_for_signal(controller.sessions_loaded) is True

    page.table.selectRow(_session_row(page, first_id))
    assert _wait_for_signal(controller.session_loaded) is True
    assert page._detail_session.id == first_id
    assert page.results_table.rowCount() == 3
    assert page.results_table.item(2, 5).text() == "否"
    assert page.results_table.item(2, 6).text() == "latency_out_of_range"
    assert "模拟结果（非真实声学测量）" in page.detail_label.text()
    assert page.notes_edit.toPlainText() == "初始备注"

    page.notes_edit.setPlainText("  历史页修改备注  ")
    page.save_notes_button.click()
    assert _wait_for_signal(controller.notes_saved) is True
    assert _wait_for_signal(controller.session_loaded) is True
    assert service.get_session(first_id).notes == "历史页修改备注"
    assert page.notes_edit.toPlainText() == "历史页修改备注"

    page.table.selectAll()
    assert set(page.selected_session_ids()) == {first_id, second_id}
    assert "已选择 2 条会话" in page.detail_label.text()
    page.sort_combo.setCurrentIndex(page.sort_combo.findData("mean_latency_asc"))
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert set(page.selected_session_ids()) == {first_id, second_id}
    assert "已选择 2 条会话" in page.detail_label.text()

    page.connection_filter.setCurrentIndex(page.connection_filter.findData("wired"))
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert page.selected_session_ids() == [second_id]
    assert _wait_for_signal(controller.session_loaded) is True
    assert page._detail_session.id == second_id


def test_history_page_delete_confirmation_and_cascade(history_context, monkeypatch):
    service, controller, page = history_context
    session_id = service.save_session(_snapshot(with_calibration=True))
    page.refresh()
    assert _wait_for_signal(controller.sessions_loaded) is True
    page.table.selectRow(_session_row(page, session_id))
    assert _wait_for_signal(controller.session_loaded) is True
    confirmation_state = {}

    def confirm(dialog):
        confirmation_state["parent"] = dialog.parent()
        confirmation_state["default_no"] = dialog.defaultButton() == dialog.button(
            QMessageBox.StandardButton.No
        )
        confirmation_state["escape_no"] = dialog.escapeButton() == dialog.button(
            QMessageBox.StandardButton.No
        )
        confirmation_state["yes_auto_default"] = dialog.button(
            QMessageBox.StandardButton.Yes
        ).autoDefault()
        confirmation_state["no_auto_default"] = dialog.button(
            QMessageBox.StandardButton.No
        ).autoDefault()
        return QMessageBox.StandardButton.Yes

    monkeypatch.setattr(QMessageBox, "exec", confirm)
    page.delete_button.click()
    assert _wait_for_signal(controller.sessions_deleted) is True
    assert _wait_for_signal(controller.sessions_loaded) is True

    assert confirmation_state == {
        "parent": page,
        "default_no": True,
        "escape_no": True,
        "yes_auto_default": False,
        "no_auto_default": True,
    }
    assert service.list_sessions() == []
    assert page.table.rowCount() == 0
    assert page.results_table.rowCount() == 0
    with sqlite3.connect(service.database.path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM test_results").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM calibration_profiles").fetchone()[0] == 0


def test_history_page_csv_and_excel_export(history_context, monkeypatch, tmp_path):
    service, controller, page = history_context
    service.save_session(_snapshot("first", "Alpha", "bluetooth"))
    service.save_session(_snapshot("second", "Beta", "wired"))
    page.refresh()
    assert _wait_for_signal(controller.sessions_loaded) is True
    page.table.selectAll()

    csv_path = tmp_path / "history-ui.csv"
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (str(csv_path), "CSV 文件 (*.csv)"),
    )
    page.export_csv_button.click()
    assert _wait_for_signal(controller.export_completed) is True
    assert csv_path.exists()

    excel_path = tmp_path / "history-ui.xlsx"
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (str(excel_path), "Excel 文件 (*.xlsx)"),
    )
    page.export_excel_button.click()
    assert _wait_for_signal(controller.export_completed) is True
    assert excel_path.exists()
    assert "已导出 2 条会话" in page.status_label.text()


def test_history_page_recovers_after_background_query_failure(history_context, monkeypatch):
    service, controller, page = history_context
    service.save_session(_snapshot())
    original_list_sessions = service.list_sessions

    def fail_list_sessions(*args, **kwargs):
        raise sqlite3.OperationalError("模拟历史查询失败")

    failures = []
    controller.operation_failed.connect(
        lambda operation, message: failures.append((operation, message))
    )
    monkeypatch.setattr(service, "list_sessions", fail_list_sessions)
    page.refresh()
    assert _wait_for_signal(controller.operation_finished) is True
    assert failures == [("load_sessions", "模拟历史查询失败")]
    assert "模拟历史查询失败" in page.status_label.text()
    assert page.refresh_button.isEnabled() is True

    monkeypatch.setattr(service, "list_sessions", original_list_sessions)
    page.refresh()
    assert _wait_for_signal(controller.sessions_loaded) is True
    assert page.table.rowCount() == 1
    assert page.status_label.text() == "已加载 1 条会话"
